"""Apply batched deco-helper catalog edits without regenerating icons.

CSV is the source of truth. This patches helper_categories.csv, the live canvas
RAW/PARENTS/CHILDREN (no ICONS rewrite), English loc in draft+live, PARENT_ORDER,
and prepends LOOK_RULES so a later generate_decoblocks.py run keeps the mapping.

  python patch_helper_catalog.py ops.json
  python patch_helper_catalog.py ops.json --dry-run
"""
from __future__ import annotations

import argparse
import ast
import csv
import json
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path
import xml.etree.ElementTree as ET

HERE = Path(__file__).resolve().parent
CSV_PATH = HERE / "helper_categories.csv"
DECOBLOCKS = HERE / "generate_decoblocks.py"
GEN_CANVAS = Path(
    r"C:\Users\rft30\.cursor\projects\c-GitHub-7D2D-Mods\agent-tools\generate_deco_helper_canvas.py"
)
CANVAS = Path(
    r"C:\Users\rft30\.cursor\projects\c-GitHub-7D2D-Mods\canvases\deco-helper-block-list.canvas.tsx"
)
CANVAS_DATA = CANVAS.with_suffix(".data.json")
HELPER_LOOK_PATH = HERE / "helper_look.csv"
KIND_SUFFIXES = (
    "Cooking",
    "SteelInsecure",
    "IronInsecure",
    "Insecure",
    "Steel",
    "Iron",
)
STATUS_TOKS = (
    "Closed",
    "Open",
    "Empty",
    "Full",
    "Working",
    "Stopped",
    "Broken",
)
COLOR_TOKS = (
    "ArmyGreen",
    "DarkGreen",
    "Aqua",
    "White",
    "Black",
    "Brown",
    "Yellow",
    "Green",
    "Grey",
    "Gray",
    "Silver",
    "Orange",
    "Purple",
    "Pink",
    "Brass",
    "Blue",
    "Red",
    "Tan",
)
CAT_STRIP_LEADING_DIGIT = re.compile(r"^\d+")
CAT_ALIASES = {
    "seats": "Seat",
    "death": "Death",
    "death1": "Death",
    "1painting": "Painting",
}
LOC_PATHS = (
    Path(r"c:\GitHub\7D2D-Mods\01_Draft\AGF-VP-DecorationBlock-v3.0.3\Config\Localization.csv"),
    Path(
        r"C:\Program Files (x86)\Steam\steamapps\common\7 Days To Die\Mods\AGF-VP-DecorationBlock-v3.0.3\Config\Localization.csv"
    ),
)
KIND_SUFFIX_RE = re.compile(r"(\s*\[[0-9a-fA-F]{6}\].*\[-\]\s*)+$")
KIND_ENGLISH_TAG = {
    "plain": "[ddcdfa](Deco)[-]",
    "deco": "[ddcdfa](Deco)[-]",
    "loot": "[c8e6be](Loot)[-]",
    "campfire": "[f0beb9](Campfire)[-]",
    "structure": "[f0cdaa](Structure)[-]",
    "electric": "[e6dca5](Electric)[-]",
    "powered": "[e6dca5](Electric)[-]",
    "light": "[e6dca5](Light)[-]",
}


def pad_sort(n: int) -> str:
    return f"{n:05d}"


def source_of(name: str) -> str:
    rest = name[7:] if name.startswith("agfDeco") else name
    for suf in KIND_SUFFIXES:
        if rest.endswith(suf):
            return rest[: -len(suf)]
    return rest


def insert_status_token(name: str, tok: str) -> list[str]:
    out = []
    for i, ch in enumerate(name):
        if i > 0 and ch.isupper():
            out.append(name[:i] + tok + name[i:])
    out.append(name + tok)
    return out


def strip_camel_token(name: str, tok: str) -> str:
    i = 0
    while True:
        at = name.find(tok, i)
        if at < 0:
            return name
        end = at + len(tok)
        before = name[at - 1] if at else ""
        after = name[end] if end < len(name) else ""
        ok_before = at == 0 or before.islower() or before.isdigit()
        ok_after = end == len(name) or after.isupper() or after.isdigit()
        if ok_before and ok_after:
            name = name[:at] + name[end:]
            i = at
            continue
        i = at + 1


def model_stem(source: str) -> str:
    """cntLockersShortClosedYellow and cntLockersShortOpenGreen -> cntLockersShort."""
    s = source
    for tok in STATUS_TOKS + COLOR_TOKS:
        s = strip_camel_token(s, tok)
    return s or source


def color_of(source: str) -> str:
    for tok in COLOR_TOKS:
        if strip_camel_token(source, tok) != source:
            return tok
    return "_"


def status_bucket(source: str, english: str = "") -> int:
    # Closed/Open labels win even when the source token is Empty/Full (armoire drawers).
    if "Closed" in source or ", Closed" in english or "(Closed" in english:
        return 0
    if "Open" in source or ", Open" in english:
        return 1
    if "Empty" in source or ", Empty" in english:
        return 2
    if "Full" in source or ", Full" in english:
        return 3
    return 4


def map_helper_to_excel(
    helper_sources: set[str],
    excel: list[dict],
) -> tuple[dict[str, float], dict[str, tuple[str, str]], set[str], int]:
    """Exact source names first, then v3 names that dropped Closed/Open from Excel."""
    excel_rank = {r["name"]: float(i) for i, r in enumerate(excel)}
    excel_cat = {r["name"]: excel_parent_child(r) for r in excel}
    helper_rank: dict[str, float] = {}
    helper_cat: dict[str, tuple[str, str]] = {}
    used_excel: set[str] = set()
    aliases = 0
    for src in helper_sources:
        if src in excel_rank:
            helper_rank[src] = excel_rank[src]
            helper_cat[src] = excel_cat[src]
            used_excel.add(src)
    for ename, rank in excel_rank.items():
        if ename in used_excel:
            continue
        for tok in STATUS_TOKS:
            if tok not in ename:
                continue
            stripped = ename.replace(tok, "", 1)
            if stripped in helper_sources and stripped not in helper_rank:
                helper_rank[stripped] = rank
                helper_cat[stripped] = excel_cat[ename]
                used_excel.add(ename)
                aliases += 1
                break
    for src in helper_sources:
        if src in helper_rank:
            continue
        done = False
        for tok in STATUS_TOKS:
            for cand in insert_status_token(src, tok):
                if cand in excel_rank and cand not in used_excel:
                    helper_rank[src] = excel_rank[cand]
                    helper_cat[src] = excel_cat[cand]
                    used_excel.add(cand)
                    aliases += 1
                    done = True
                    break
            if done:
                break
    for src in helper_sources:
        if src in helper_rank:
            continue
        best = ""
        for ename in excel_rank:
            if src.startswith(ename) and len(ename) >= 8 and len(ename) > len(best):
                best = ename
        if best:
            helper_rank[src] = excel_rank[best] + 0.001
            helper_cat[src] = excel_cat[best]
            aliases += 1
    return helper_rank, helper_cat, used_excel, aliases


def nice_label(raw: object) -> str:
    s = str(raw or "").strip()
    if not s:
        return ""
    key = s.lower()
    if key in CAT_ALIASES:
        return CAT_ALIASES[key]
    stripped = CAT_STRIP_LEADING_DIGIT.sub("", s)
    if stripped:
        s = stripped
        key = s.lower()
        if key in CAT_ALIASES:
            return CAT_ALIASES[key]
    if s.islower() or s.isupper():
        return s.replace("_", " ").title()
    return s


def excel_parent_child(row: dict) -> tuple[str, str]:
    """Canvas has two levels. Excel C3 is the real subcategory (camping vs couch)."""
    c1 = nice_label(row.get("c1"))
    if not c1:
        return "", ""

    def child_part(raw: object) -> str:
        lab = nice_label(raw)
        if not lab or lab.isdigit():
            return ""
        return lab

    child = child_part(row.get("c3")) or child_part(row.get("c2")) or child_part(row.get("c4")) or "General"
    return c1, child


def load_excel_support(path: Path) -> list[dict]:
    import openpyxl

    src = path
    if not src.is_file():
        raise SystemExit(f"excel not found: {src}")
    try:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    except PermissionError:
        copy = HERE / "_tmp_excel_support.xlsx"
        shutil.copy2(src, copy)
        src = copy
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = []
    seen = set()
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        name = str(row[0] or "").strip()
        if not name or name in seen:
            continue
        seen.add(name)
        out.append(
            {
                "name": name,
                "c1": row[2] if len(row) > 2 else "",
                "c2": row[3] if len(row) > 3 else "",
                "c3": row[4] if len(row) > 4 else "",
                "c4": row[5] if len(row) > 5 else "",
            }
        )
    wb.close()
    return out


def merge_canvas_edits(rows: list[dict[str, str]]) -> None:
    if not CANVAS_DATA.is_file():
        return
    data = json.loads(CANVAS_DATA.read_text(encoding="utf-8"))
    edits = data.get("edits_v1") or {}
    for r in rows:
        e = edits.get(r["BlockName"]) or {}
        if e.get("sort"):
            r["SortOrder1"] = str(e["sort"])
        if e.get("cat1"):
            r["Parent"] = str(e["cat1"])
        if e.get("cat2"):
            r["Child"] = str(e["cat2"])
    rows.sort(key=lambda r: (r["SortOrder1"], r["BlockName"]))


def write_canvas_edits(rows: list[dict[str, str]], dry: bool) -> None:
    data = {}
    if CANVAS_DATA.is_file():
        data = json.loads(CANVAS_DATA.read_text(encoding="utf-8"))
    edits = {}
    for r in rows:
        edits[r["BlockName"]] = {
            "sort": r["SortOrder1"],
            "cat1": r["Parent"],
            "cat2": r["Child"],
        }
    data["edits_v1"] = edits
    if dry:
        print(f"canvas data would write {len(edits)} edits_v1")
        return
    CANVAS_DATA.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {CANVAS_DATA.name} edits_v1={len(edits)}")


def write_helper_look(rows: list[dict[str, str]], dry: bool) -> None:
    by_source: dict[str, tuple[str, str]] = {}
    for r in rows:
        if r.get("Kind") == "campfire":
            continue
        src = source_of(r["BlockName"])
        by_source.setdefault(src, (r["Parent"], r["Child"]))
    if dry:
        print(f"helper_look would write {len(by_source)} sources")
        return
    with HELPER_LOOK_PATH.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Source", "Parent", "Child"])
        w.writeheader()
        for src, (p, c) in sorted(by_source.items()):
            w.writerow({"Source": src, "Parent": p, "Child": c})
    print(f"wrote {HELPER_LOOK_PATH.name} ({len(by_source)} sources)")


def set_parent_order(parents: list[str], dry: bool) -> None:
    text = GEN_CANVAS.read_text(encoding="utf-8")
    m = re.search(r"PARENT_ORDER = (\[[\s\S]*?\n\])", text)
    if not m:
        raise SystemExit("PARENT_ORDER missing")
    blob = "[\n    " + ",\n    ".join(f'"{p}"' for p in parents) + ",\n]"
    text = text[: m.start(1)] + blob + text[m.end(1) :]
    if dry:
        print(f"PARENT_ORDER would set {len(parents)} parents")
        return
    GEN_CANVAS.write_text(text, encoding="utf-8")
    print(f"PARENT_ORDER = {parents}")


def apply_excel_support(
    rows: list[dict[str, str]],
    op: dict,
    dry: bool,
) -> tuple[list[str], dict[str, list[str]]]:
    """Excel row order + Category/Category2 for helper items that already exist.

    Does not add Excel-only blocks (doors, missing models). Campfire twins stay
    first as Kitchen/Campfires. v3-only helper items keep current cats and sit
    after same-family Excel items.
    """
    xlsx = Path(op["xlsx"])
    excel = load_excel_support(xlsx)
    merge_canvas_edits(rows)

    helper_sources = {source_of(r["BlockName"]) for r in rows}
    helper_rank, helper_cat, used_excel, aliases = map_helper_to_excel(
        helper_sources, excel
    )
    skipped = [r["name"] for r in excel if r["name"] not in used_excel]
    v3_only = sorted(helper_sources - set(helper_rank))

    camp = [r for r in rows if r.get("Kind") == "campfire"]
    rest = [r for r in rows if r.get("Kind") != "campfire"]

    for i, r in enumerate(camp):
        r["Parent"] = "Kitchen"
        r["Child"] = "Campfires"
        r["_rank"] = float(i)

    family_last: dict[str, float] = {}
    family_status_last: dict[str, dict[int, float]] = defaultdict(dict)
    family_v3_i: dict[tuple[str, int], int] = defaultdict(int)
    family_cats: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for r in rest:
        src = source_of(r["BlockName"])
        if src in helper_rank:
            rank = helper_rank[src]
            fam = r["HelperFamily"]
            family_last[fam] = max(family_last.get(fam, -1.0), rank)
            b = status_bucket(src, r.get("EnglishName") or "")
            prev = family_status_last[fam].get(b)
            family_status_last[fam][b] = rank if prev is None else max(prev, rank)
        p, c = helper_cat.get(src, ("", ""))
        if p:
            family_cats[r["HelperFamily"]].append((p, c))

    def family_cat(fam: str) -> tuple[str, str] | None:
        counts: dict[tuple[str, str], int] = {}
        for pair in family_cats.get(fam) or []:
            counts[pair] = counts.get(pair, 0) + 1
        if not counts:
            return None
        return max(counts.items(), key=lambda kv: kv[1])[0]

    recat = 0
    inherited = 0
    for r in rest:
        src = source_of(r["BlockName"])
        p, c = helper_cat.get(src, ("", ""))
        if p:
            r["Parent"] = p
            r["Child"] = c
            recat += 1
        else:
            pick = family_cat(r["HelperFamily"])
            if pick:
                r["Parent"], r["Child"] = pick
                inherited += 1
        if src in helper_rank:
            r["_rank"] = 1000.0 + helper_rank[src]
        else:
            fam = r["HelperFamily"]
            b = status_bucket(src, r.get("EnglishName") or "")
            family_v3_i[(fam, b)] += 1
            base = family_status_last[fam].get(b)
            if base is None:
                base = family_last.get(fam, 40_000.0)
            r["_rank"] = 1000.0 + base + 0.01 * family_v3_i[(fam, b)]

    by_stem: dict[str, list[dict[str, str]]] = defaultdict(list)
    for r in rest:
        by_stem[model_stem(source_of(r["BlockName"]))].append(r)
    for members in by_stem.values():
        buckets = {
            status_bucket(source_of(r["BlockName"]), r.get("EnglishName") or "")
            for r in members
        }
        if 0 in buckets and 1 in buckets:
            anchor = min(r["_rank"] for r in members)
            for r in members:
                src = source_of(r["BlockName"])
                b = status_bucket(src, r.get("EnglishName") or "")
                r["_rank"] = anchor + (b * 0.5) + ((r["_rank"] - anchor) * 1e-6)
        if 2 in buckets and 3 in buckets:
            empty_full = [
                r
                for r in members
                if status_bucket(source_of(r["BlockName"]), r.get("EnglishName") or "") in (2, 3)
            ]
            anchor = min(r["_rank"] for r in empty_full)
            for r in empty_full:
                src = source_of(r["BlockName"])
                b = status_bucket(src, r.get("EnglishName") or "")
                r["_rank"] = anchor + ((b - 2) * 0.5) + ((r["_rank"] - anchor) * 1e-6)
        closed = [r for r in members if status_bucket(source_of(r["BlockName"]), r.get("EnglishName") or "") == 0]
        opened = [r for r in members if status_bucket(source_of(r["BlockName"]), r.get("EnglishName") or "") == 1]
        if closed and opened:
            color_order = []
            for r in sorted(closed, key=lambda x: x["_rank"]):
                col = color_of(source_of(r["BlockName"]))
                if col not in color_order:
                    color_order.append(col)
            if color_order:
                open_anchor = max(r["_rank"] for r in closed)
                order = list(color_order)

                def open_key(row: dict[str, str], order: list[str] = order) -> tuple:
                    col = color_of(source_of(row["BlockName"]))
                    idx = order.index(col) if col in order else 99
                    return (idx, row["_rank"])

                for i, r in enumerate(sorted(opened, key=open_key)):
                    r["_rank"] = open_anchor + 0.001 * (i + 1)

    rest.sort(key=lambda r: (r["_rank"], r["BlockName"]))
    ordered = camp + rest
    for i, r in enumerate(ordered, 1):
        r["SortOrder1"] = pad_sort(i)
        r.pop("_rank", None)
    rows[:] = ordered

    parents: list[str] = []
    extra: dict[str, list[str]] = {}
    if camp:
        parents.append("Kitchen")
        extra["Kitchen"] = ["Campfires"]
    for r in rows:
        p, c = r["Parent"], r["Child"]
        if p not in parents:
            parents.append(p)
        extra.setdefault(p, [])
        if c not in extra[p]:
            extra[p].append(c)

    set_parent_order(parents, dry)
    write_helper_look(rows, dry)
    write_canvas_edits(rows, dry)
    print(
        f"apply_excel_support helper={len(rows)} recat={recat} inherited={inherited} "
        f"aliases={aliases} campfire={len(camp)} excel-skipped={len(skipped)} "
        f"v3-only-sources={len(v3_only)} parents={len(parents)}"
    )
    print(f"  parents: {parents}")
    print(f"  excel names not on helper (not added): {len(skipped)}")
    print(f"  v3-only sources kept: {len(v3_only)}")
    return parents, extra


def load_csv() -> tuple[list[dict[str, str]], list[str]]:
    with CSV_PATH.open(encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
        fields = list(rows[0].keys())
    return rows, fields


def save_csv(rows: list[dict[str, str]], fields: list[str]) -> None:
    with CSV_PATH.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def by_id(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {r["BlockName"]: r for r in rows}


def filtered_view(
    rows: list[dict[str, str]],
    parent: str | None,
    child: str | None,
) -> list[dict[str, str]]:
    out = []
    for r in rows:
        if parent and r["Parent"] != parent:
            continue
        if child and r["Child"] != child:
            continue
        out.append(r)
    out.sort(key=lambda r: (r["SortOrder1"], r["BlockName"]))
    return out


def select_rows(rows: list[dict[str, str]], op: dict) -> list[dict[str, str]]:
    if op.get("ids"):
        idx = by_id(rows)
        missing = [i for i in op["ids"] if i not in idx]
        if missing:
            raise SystemExit(f"unknown ids: {missing[:8]}")
        return [idx[i] for i in op["ids"]]
    parent = op.get("filter_parent")
    child = op.get("filter_child")
    view = filtered_view(rows, parent, child)
    if "from" in op or "to" in op:
        a = int(op.get("from", 1))
        b = int(op.get("to", len(view)))
        if a < 1 or b < a or b > len(view):
            raise SystemExit(
                f"range {a}-{b} out of 1..{len(view)} "
                f"(filter_parent={parent!r} filter_child={child!r})"
            )
        return view[a - 1 : b]
    if op.get("family"):
        needle = op["family"].lower()
        return [r for r in view if needle in r["HelperFamily"].lower()]
    if op.get("name_contains"):
        needle = op["name_contains"].lower()
        return [r for r in view if needle in r["EnglishName"].lower() or needle in r["BlockName"].lower()]
    raise SystemExit(f"op needs ids, from/to, family, or name_contains: {op}")


def extract_js_value(text: str, name: str, start: str) -> tuple[object, int, int]:
    if name == "RAW":
        j = text.rfind("RAW = [")
        if j < 0:
            raise SystemExit("canvas missing RAW = [")
        j = text.find("[", j)
        obj, n = json.JSONDecoder().raw_decode(text[j:])
        return obj, j, j + n
    patterns = (f"const {name}:", f"let {name}:", f"{name} =")
    i = -1
    for p in patterns:
        i = text.find(p)
        if i >= 0:
            break
    if i < 0:
        raise SystemExit(f"canvas missing {name}")
    eq = text.find("=", i)
    if eq < 0:
        raise SystemExit(f"canvas {name} has no =")
    j = text.find(start, eq)
    if j < 0:
        raise SystemExit(f"canvas {name} missing {start}")
    obj, n = json.JSONDecoder().raw_decode(text[j:])
    return obj, j, j + n


def splice(text: str, lo: int, hi: int, value: object) -> str:
    blob = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return text[:lo] + blob + text[hi:]


def patch_canvas(
    rows: list[dict[str, str]],
    parents: list[str],
    children: dict[str, list[str]],
    dry: bool,
    raw_inserts: dict[str, tuple[str, str]] | None = None,
) -> None:
    text = CANVAS.read_text(encoding="utf-8")
    pobj, plo, phi = extract_js_value(text, "PARENTS", "[")
    text = splice(text, plo, phi, parents)
    cobj, clo, chi = extract_js_value(text, "CHILDREN", "{")
    text = splice(text, clo, chi, children)
    raw, rlo, rhi = extract_js_value(text, "RAW", "[")
    keep = {r["BlockName"] for r in rows}
    raw = [cell for cell in raw if cell and cell[0] in keep]
    idx = {row[0]: i for i, row in enumerate(raw)}
    miss = 0
    raw_inserts = raw_inserts or {}
    for r in rows:
        i = idx.get(r["BlockName"])
        if i is None:
            spec = raw_inserts.get(r["BlockName"])
            after_id, icon_from = spec if spec else ("", "")
            ai = idx.get(after_id)
            ii = idx.get(icon_from) if icon_from else ai
            if ai is None or ii is None:
                miss += 1
                continue
            cell = list(raw[ii])
            cell[0] = r["BlockName"]
            cell[1] = r["EnglishName"]
            cell[2] = r["SortOrder1"]
            cell[3] = r["Parent"]
            cell[4] = r["Child"]
            raw.insert(ai + 1, cell)
            idx = {row[0]: n for n, row in enumerate(raw)}
            continue
        cell = raw[i]
        # id, name, sort, cat1, cat2, icon, tint
        cell[1] = r["EnglishName"]
        cell[2] = r["SortOrder1"]
        cell[3] = r["Parent"]
        cell[4] = r["Child"]
    text = splice(text, rlo, rhi, raw)
    if miss:
        print(f"canvas RAW miss {miss} ids (csv-only)")
    if dry:
        print(f"canvas would write {CANVAS} ({len(text)} chars)")
        return
    CANVAS.write_text(text, encoding="utf-8")
    print(f"wrote {CANVAS.name}")


def ensure_parent_order(name: str, after: str | None, dry: bool) -> None:
    text = GEN_CANVAS.read_text(encoding="utf-8")
    m = re.search(r"PARENT_ORDER = (\[[\s\S]*?\n\])", text)
    if not m:
        raise SystemExit("PARENT_ORDER missing")
    order = list(ast.literal_eval(m.group(1)))
    if name in order:
        return
    if after and after in order:
        order.insert(order.index(after) + 1, name)
    else:
        order.append(name)
    blob = "[\n    " + ",\n    ".join(f'"{p}"' for p in order) + ",\n]"
    text = text[: m.start(1)] + blob + text[m.end(1) :]
    if dry:
        print(f"PARENT_ORDER would add {name}")
        return
    GEN_CANVAS.write_text(text, encoding="utf-8")
    print(f"PARENT_ORDER += {name}")


def prepend_look_rule(needles: list[str], parent: str, child: str, dry: bool) -> None:
    if not needles:
        return
    text = DECOBLOCKS.read_text(encoding="utf-8")
    marker = "LOOK_RULES: tuple[tuple[tuple[str, ...], str, str], ...] = ("
    i = text.find(marker)
    if i < 0:
        raise SystemExit("LOOK_RULES missing")
    tup = "(" + ", ".join(json.dumps(n) for n in needles) + f'), "{parent}", "{child}"'
    insert = f"{marker}\n    ({tup}),"
    if tup in text:
        print("LOOK_RULES already has this mapping")
        return
    text = text[:i] + insert + text[i + len(marker) :]
    if dry:
        print(f"LOOK_RULES would prepend {needles} -> {parent}/{child}")
        return
    DECOBLOCKS.write_text(text, encoding="utf-8")
    print(f"LOOK_RULES += {needles} -> {parent}/{child}")


def _xml_block_snippet(els: list[ET.Element]) -> str:
    wrap = ET.Element("wrap")
    for el in els:
        wrap.append(el)
    ET.indent(wrap, space="\t")
    chunks = []
    for child in list(wrap):
        raw = ET.tostring(child, encoding="unicode")
        lines = []
        for line in raw.splitlines():
            lines.append("\t" + line if line else line)
        chunks.append("\n".join(lines))
    return "\n".join(chunks)


def _inject_xml_after(xml: str, after_name: str, snippet: str) -> str:
    needle = f'<block name="{after_name}"'
    last = xml.rfind(needle)
    if last < 0:
        raise SystemExit(f"blocks.xml missing {after_name}")
    end = xml.find("</block>", last)
    if end < 0:
        raise SystemExit(f"blocks.xml unclosed {after_name}")
    end += len("</block>")
    return xml[:end] + "\n" + snippet.rstrip() + xml[end:]


def _insert_place_alt(xml: str, after_name: str, new_name: str) -> str:
    token = f"{after_name},"
    insert = f"{after_name},{new_name},"
    if insert in xml:
        return xml
    n = xml.count(token)
    if n != 1:
        # Prefer the picker token that is not Iron/Steel: after_name + comma.
        if n == 0:
            token = f"{after_name}\""
            if token not in xml:
                raise SystemExit(f"PlaceAlt missing {after_name} ({n})")
            return xml.replace(token, f"{after_name},{new_name}\"", 1)
        raise SystemExit(f"PlaceAlt token {after_name}, count={n}")
    return xml.replace(token, insert, 1)


def _append_helper_look(source: str, parent: str, child: str, dry: bool) -> None:
    if not HELPER_LOOK_PATH.is_file():
        return
    raw = HELPER_LOOK_PATH.read_text(encoding="utf-8").splitlines()
    header, lines = raw[0], raw[1:]
    sources = {ln.split(",", 1)[0] for ln in lines if ln.strip()}
    if source in sources:
        return
    row = f"{source},{parent},{child}"
    lines.append(row)
    body = [header]
    body.extend(sorted((ln for ln in lines if ln.strip()), key=lambda s: s.split(",", 1)[0].lower()))
    if dry:
        print(f"helper_look would add {source}")
        return
    HELPER_LOOK_PATH.write_text("\n".join(body) + "\n", encoding="utf-8")
    print(f"helper_look += {source}")


def _append_loc_rows(new_rows: list[dict[str, str]], dry: bool) -> None:
    if not new_rows:
        return
    import generate_decoblocks as g

    blobs = [g.format_loc_row(r) for r in new_rows]
    keys = {r["Key"] for r in new_rows}
    for path in LOC_PATHS:
        if not path.exists():
            print(f"skip loc {path} (missing)")
            continue
        raw = path.read_text(encoding="utf-8").splitlines()
        have = {ln.split(",", 1)[0] for ln in raw[1:] if ln.strip()}
        add = [line for row, line in zip(new_rows, blobs) if row["Key"] not in have]
        if dry:
            print(f"loc {path.name} would append {len(add)}")
            continue
        if not add:
            print(f"loc {path.name} already has {len(keys)} keys")
            continue
        text = "\n".join(raw + add) + "\n"
        path.write_text(text, encoding="utf-8")
        print(f"loc {path.name} appended {len(add)}")


def add_canvas_edit_keys(new_rows: list[dict[str, str]], dry: bool) -> None:
    if not new_rows or not CANVAS_DATA.is_file():
        return
    data = json.loads(CANVAS_DATA.read_text(encoding="utf-8"))
    edits = data.get("edits_v1") or {}
    for r in new_rows:
        edits[r["BlockName"]] = {
            "sort": r["SortOrder1"],
            "cat1": r["Parent"],
            "cat2": r["Child"],
        }
    data["edits_v1"] = edits
    if dry:
        print(f"canvas data would add {len(new_rows)} edits_v1")
        return
    CANVAS_DATA.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"wrote {CANVAS_DATA.name} edits_v1={len(edits)}")


def add_helper_sources(
    rows: list[dict[str, str]],
    op: dict,
    loc_updates: dict[str, str],
    raw_inserts: dict[str, tuple[str, str]],
    dry: bool,
) -> None:
    """Emit loot/upgradeable families for listed vanilla sources and inject them."""
    import generate_decoblocks as g

    specs = op.get("sources") or []
    if not specs:
        raise SystemExit("add_helper_sources needs sources[]")
    flat_map, _order, _all_names = g.load_flat_map()
    vanilla = g.load_loc_map(g.GAME_LOC)
    idx = {r["BlockName"]: i for i, r in enumerate(rows)}
    loc_rows: list[dict[str, str]] = []
    xml_jobs: list[tuple[str, list[ET.Element]]] = []
    picker_jobs: list[tuple[str, str]] = []
    new_catalog: list[dict[str, str]] = []

    for spec in specs:
        source = spec["source"]
        emit_kind = spec["kind"]
        after_id = spec["after_id"]
        english = spec["english"]
        parent = spec["parent"]
        child = spec["child"]
        xml_after = spec["xml_after"]
        icon_from = spec.get("icon_from") or after_id
        picker_id = g.deco_name(source)
        if picker_id in idx:
            print(f"skip existing {picker_id}")
            continue
        flat = flat_map.get(source)
        if flat is None:
            raise SystemExit(f"unknown vanilla source {source}")
        after = next((r for r in rows if r["BlockName"] == after_id), None)
        if after is None:
            raise SystemExit(f"after_id missing {after_id}")
        sort1 = after["SortOrder1"]
        family, _picker = g.emit_kind(
            emit_kind, source, flat, sort1, {}, include_cooking=False
        )
        xml_jobs.append((xml_after, family))
        picker_jobs.append((after_id, picker_id))
        kind_col = "loot" if emit_kind == "loot" else "structure"
        kind_tag = "loot" if emit_kind == "loot" else "structure"
        helper_family = after["HelperFamily"] if kind_col == "structure" else source
        cat_row = {
            "BlockName": picker_id,
            "EnglishName": apply_kind_tag(english, kind_tag),
            "SortOrder1": sort1,
            "SortOrder2": after.get("SortOrder2") or "",
            "Kind": kind_col,
            "HelperFamily": helper_family,
            "Parent": parent,
            "Child": child,
        }
        insert_at = idx[after_id] + 1
        rows.insert(insert_at, cat_row)
        idx = {r["BlockName"]: i for i, r in enumerate(rows)}
        loc_updates[picker_id] = cat_row["EnglishName"]
        raw_inserts[picker_id] = (after_id, icon_from)
        new_catalog.append(cat_row)
        _append_helper_look(source, parent, child, dry)

        name_set = {el.get("name") or "" for el in family}
        for el in family:
            name = el.get("name") or ""
            src, fam, tier = g.split_deco_name(name, name_set)
            vanilla_row = vanilla.get(src, {})
            lang_names: dict[str, str] = {}
            bases: dict[str, str] = {}
            for lang in g.LOC_LANGS:
                base = ""
                vanilla_text = vanilla_row.get(lang) or ""
                if vanilla_text and not g.loc_text_broken(vanilla_text):
                    base = vanilla_text
                if not base and lang != "english":
                    base = bases.get("english", "")
                if not base:
                    base = g.humanize_source(src)
                bases[lang] = base
                lang_names[lang] = g.compose_loc_name(base, fam, tier, lang, src)
            loc_rows.append(g.loc_row(name, lang_names))

    if dry:
        print(f"add_helper_sources would add {len(new_catalog)} picker rows")
        return

    snippet_by_after: dict[str, list[str]] = {}
    for xml_after, family in xml_jobs:
        snippet_by_after.setdefault(xml_after, []).append(_xml_block_snippet(family))
    for path in g.OUTPUT_BLOCKS:
        if not path.exists():
            print(f"skip blocks {path} (missing)")
            continue
        xml = path.read_text(encoding="utf-8")
        for xml_after, parts in snippet_by_after.items():
            xml = _inject_xml_after(xml, xml_after, "\n".join(parts))
        for after_id, picker_id in picker_jobs:
            xml = _insert_place_alt(xml, after_id, picker_id)
        path.write_text(xml, encoding="utf-8")
        print(f"injected {path}")
    _append_loc_rows(loc_rows, dry=False)
    add_canvas_edit_keys(new_catalog, dry=False)
    print(f"add_helper_sources added {len(new_catalog)} picker rows")


def apply_kind_tag(english: str, kind: str) -> str:
    tag = KIND_ENGLISH_TAG.get(kind)
    if not tag:
        raise SystemExit(f"unknown kind {kind!r}")
    base = KIND_SUFFIX_RE.sub("", english).rstrip()
    return f"{base} {tag}"


def rename_english(row: dict[str, str], new: str) -> str:
    old = row["EnglishName"]
    m = KIND_SUFFIX_RE.search(old)
    suffix = m.group(0) if m else ""
    if KIND_SUFFIX_RE.search(new):
        return new
    return (new.rstrip() + suffix).strip()


def _csv_line(fields: list[str]) -> str:
    out = []
    for i, val in enumerate(fields):
        if i == 7:
            out.append("")
            continue
        if i >= 6:
            out.append('"' + val.replace('"', '""') + '"')
        else:
            out.append(val)
    return ",".join(out)


def patch_loc_english(keys: dict[str, str], dry: bool) -> None:
    if not keys:
        return
    for path in LOC_PATHS:
        if not path.exists():
            print(f"skip loc {path} (missing)")
            continue
        raw = path.read_text(encoding="utf-8").splitlines()
        n = 0
        new_lines = [raw[0]]
        for line in raw[1:]:
            if not line.strip():
                new_lines.append(line)
                continue
            key = line.split(",", 1)[0]
            if key not in keys:
                new_lines.append(line)
                continue
            row = next(csv.reader([line]))
            row[6] = keys[key]
            new_lines.append(_csv_line(row))
            n += 1
        if dry:
            print(f"loc {path.name} would update {n}")
            continue
        path.write_text("\n".join(new_lines) + "\n", encoding="utf-8")
        print(f"loc {path.name} updated {n}")


def rebuild_children(rows: list[dict[str, str]], parents: list[str], extra: dict[str, list[str]]) -> dict[str, list[str]]:
    seen: dict[str, list[str]] = defaultdict(list)
    have: dict[str, set[str]] = defaultdict(set)
    for r in rows:
        p, c = r["Parent"], r["Child"]
        if c not in have[p]:
            have[p].add(c)
            seen[p].append(c)
    out: dict[str, list[str]] = {}
    for p in parents:
        kids = list(extra.get(p, []))
        for c in seen.get(p, []):
            if c not in kids:
                kids.append(c)
        out[p] = kids
    return out


def reorder_by_category(rows: list[dict[str, str]], parents: list[str], children: dict[str, list[str]]) -> None:
    def rank(r: dict[str, str]) -> tuple:
        p = r["Parent"]
        try:
            pi = parents.index(p)
        except ValueError:
            pi = len(parents)
        kids = children.get(p, [])
        try:
            ci = kids.index(r["Child"])
        except ValueError:
            ci = len(kids)
        return (pi, p, ci, r["Child"], r["SortOrder1"], r["BlockName"])

    rows.sort(key=rank)
    for i, r in enumerate(rows, 1):
        r["SortOrder1"] = pad_sort(i)


def bake_ingame(rows: list[dict[str, str]], dry: bool) -> None:
    """Stamp XML SortOrder1 + PlaceAlt so the in-game picker matches catalog order."""
    draft = Path(
        r"c:\GitHub\7D2D-Mods\01_Draft\AGF-VP-DecorationBlock-v3.0.3\Config\blocks.xml"
    )
    live = Path(
        r"C:\Program Files (x86)\Steam\steamapps\common\7 Days To Die\Mods\AGF-VP-DecorationBlock-v3.0.3\Config\blocks.xml"
    )
    stamps: dict[str, str] = {}
    family_suf = ("Iron", "Steel", "Insecure", "IronInsecure", "SteelInsecure")
    for r in rows:
        sort = r["SortOrder1"]
        base = r["BlockName"]
        stamps[base] = sort
        if not base.endswith("Cooking"):
            for suf in family_suf:
                stamps[base + suf] = sort
    picker = ",".join(r["BlockName"] for r in rows)
    if dry:
        print(f"bake would stamp {len(stamps)} xml names, PlaceAlt {len(rows)}")
        return
    text = draft.read_text(encoding="utf-8")
    so1_re = re.compile(r'(<property name="SortOrder1" value=")[^"]+(" />)')

    def repl_block(m: re.Match[str]) -> str:
        name = m.group(1)
        sort = stamps.get(name)
        if not sort:
            return m.group(0)
        return so1_re.sub(rf"\g<1>{sort}\g<2>", m.group(0), count=1)

    text, nblock = re.subn(
        r'<block name="(agfDeco[^"]+)">[\s\S]*?</block>',
        repl_block,
        text,
    )
    text, nalt = re.subn(
        r'(<property name="PlaceAltBlockValue" value=")[^"]*(" />)',
        lambda m: m.group(1) + picker + m.group(2),
        text,
        count=1,
    )
    draft.write_text(text, encoding="utf-8")
    shutil.copy2(draft, live)
    print(f"baked xml SortOrder1 on {nblock} blocks, PlaceAlt replaced={nalt}, copied live")


def apply_ops(ops: list[dict], dry: bool) -> None:
    rows, fields = load_csv()
    merge_canvas_edits(rows)
    text = CANVAS.read_text(encoding="utf-8")
    parents, _, _ = extract_js_value(text, "PARENTS", "[")
    children, _, _ = extract_js_value(text, "CHILDREN", "{")
    parents = list(parents)
    children = {k: list(v) for k, v in children.items()}
    extra_children: dict[str, list[str]] = {k: list(v) for k, v in children.items()}
    loc_updates: dict[str, str] = {}
    raw_inserts: dict[str, tuple[str, str]] = {}
    do_reorder = False
    do_write_edits = False
    do_bake = False
    touched: list[str] = []

    for op in ops:
        kind = op["op"]
        if kind == "add_parent":
            name = op["name"]
            if name not in parents:
                after = op.get("after")
                if after and after in parents:
                    parents.insert(parents.index(after) + 1, name)
                else:
                    parents.append(name)
            kids = op.get("children") or []
            extra_children.setdefault(name, [])
            for c in kids:
                if c not in extra_children[name]:
                    extra_children[name].append(c)
            ensure_parent_order(name, op.get("after"), dry)
            print(f"parent {name} children={extra_children[name]}")
        elif kind == "add_child":
            p = op["parent"]
            c = op["name"]
            extra_children.setdefault(p, [])
            if c not in extra_children[p]:
                extra_children[p].append(c)
            if p not in parents:
                parents.append(p)
            print(f"child {p}/{c}")
        elif kind == "set_cat":
            target_p = op["parent"]
            target_c = op["child"]
            extra_children.setdefault(target_p, [])
            if target_c not in extra_children[target_p]:
                extra_children[target_p].append(target_c)
            if target_p not in parents:
                raise SystemExit(f"unknown parent {target_p!r}; add_parent first")
            picked = select_rows(rows, op)
            for r in picked:
                r["Parent"] = target_p
                r["Child"] = target_c
                touched.append(r["BlockName"])
            print(f"set_cat {len(picked)} -> {target_p}/{target_c}")
            if op.get("needles"):
                prepend_look_rule(list(op["needles"]), target_p, target_c, dry)
        elif kind == "rename":
            picked = select_rows(rows, op)
            if len(picked) != 1 and not op.get("english"):
                raise SystemExit("rename needs one row or shared english")
            new = op["english"]
            for r in picked:
                r["EnglishName"] = rename_english(r, new)
                loc_updates[r["BlockName"]] = r["EnglishName"]
                touched.append(r["BlockName"])
            print(f"rename {len(picked)} -> {new}")
        elif kind == "set_kind":
            target = op["kind"]
            if target not in KIND_ENGLISH_TAG:
                raise SystemExit(f"unknown kind {target!r}")
            picked = select_rows(rows, op)
            for r in picked:
                r["Kind"] = "plain" if target == "deco" else target
                r["EnglishName"] = apply_kind_tag(r["EnglishName"], target)
                loc_updates[r["BlockName"]] = r["EnglishName"]
                touched.append(r["BlockName"])
            print(f"set_kind {len(picked)} -> {target}")
        elif kind == "reorder_by_category":
            do_reorder = True
        elif kind == "apply_excel_support":
            parents, extra_children = apply_excel_support(rows, op, dry)
        elif kind == "apply_helper_review":
            from helper_review import apply_helper_review as _review

            parents, extra_children, loc_part = _review(rows, sys.modules[__name__], dry)
            loc_updates.update(loc_part)
        elif kind == "add_helper_sources":
            add_helper_sources(rows, op, loc_updates, raw_inserts, dry)
        elif kind == "apply_list_punch":
            from helper_review import apply_list_punch as _punch

            loc_updates.update(_punch(rows, sys.modules[__name__], dry))
            do_write_edits = True
        elif kind == "bake_ingame":
            do_bake = True
        else:
            raise SystemExit(f"unknown op {kind}")

    children = rebuild_children(rows, parents, extra_children)
    if do_reorder:
        print(
            "skip reorder_by_category: that would rewrite SortOrder1; "
            "category grouping is a canvas view only"
        )

    if dry:
        print(f"csv would write {len(rows)} rows; loc keys {len(loc_updates)}")
        patch_canvas(rows, parents, children, dry=True, raw_inserts=raw_inserts)
        if do_write_edits:
            write_canvas_edits(rows, dry=True)
        if do_bake:
            bake_ingame(rows, dry=True)
        return
    save_csv(rows, fields)
    print(f"wrote {CSV_PATH.name} ({len(rows)} rows)")
    patch_loc_english(loc_updates, dry=False)
    patch_canvas(rows, parents, children, dry=False, raw_inserts=raw_inserts)
    if do_write_edits:
        write_canvas_edits(rows, dry=False)
        write_helper_look(rows, dry=False)
    if do_bake:
        bake_ingame(rows, dry=False)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("ops", type=Path, help="JSON file with {ops: [...]}")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    payload = json.loads(args.ops.read_text(encoding="utf-8"))
    ops = payload["ops"] if isinstance(payload, dict) else payload
    apply_ops(ops, dry=args.dry_run)


if __name__ == "__main__":
    main()
